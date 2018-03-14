# -*- coding:utf-8 -*-

# create:2018-03-09
# author:zengln
# desc: 将 word 文档中的接口转化为 xml

import docx
import re
from lxml import etree


#  判断传入的根节点中是否包含子节点 ,
#  子节点存在返回 True , 子节点不存在返回 False
def isCludeNode(root, nodeName):
    isExist = root.find(nodeName.strip())

    if isExist is None:
        return False

    return True


file = docx.Document("C:\\Users\Desktop\\test.docx")
xmlfile = open("test.xml", "w")
for table in file.tables:

    if table.cell(0, 0).text != "节点名":
        continue

    row_count = len(table.rows)
    tree1 = etree.Element("agent")

    for row in range(1, row_count):
        nodeStr = table.cell(row, 0).text
        nodeList = nodeStr.split("/")

        for i, node in enumerate(nodeList):
            result = re.search(r'(.*?)\|(.*?)', node, re.S)

            if not result:
                   continue

            nodeList[i] = result.group(1)

        for i in range(2, len(nodeList)):
            if not isCludeNode(locals()['tree' + str(i-1)], nodeList[i]):
                locals()['tree' + str(i)] = etree.SubElement(locals()['tree' + str(i - 1)], nodeList[i].strip())

    print(etree.tostring(tree1))
    xmlfile.write(etree.tostring(tree1).decode('utf-8'))
    xmlfile.write("\r\n")

xmlfile.close()